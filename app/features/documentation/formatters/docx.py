from __future__ import annotations

import re
from io import BytesIO
from datetime import UTC, datetime

from docx import Document as WordDocument
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

from app.models.document import Document


_CODE_FENCE_RE = re.compile(r"```(?:\w+)?\n(.*?)```", re.DOTALL)
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_ITALIC_RE = re.compile(r"\*(.+?)\*")


class DocxFormatter:
    def format(self, doc: Document) -> bytes:
        word = WordDocument()

        # Document title
        word.add_heading(doc.title, level=0)
        word.add_paragraph(
            f"Type: {doc.doc_type}  |  Version: {doc.version}  |  "
            f"Generated: {datetime.now(UTC).strftime('%Y-%m-%d')}",
            style="Intense Quote",
        )
        word.add_paragraph()

        if not doc.content_markdown:
            word.add_paragraph("No content generated.")
            return self._to_bytes(word)

        self._render_markdown(word, doc.content_markdown)
        return self._to_bytes(word)

    def _render_markdown(self, word: WordDocument, content: str) -> None:
        # Split by code fences first to preserve them
        parts = _CODE_FENCE_RE.split(content)
        # parts alternates: [text, code_block, text, code_block, ...]
        for i, part in enumerate(parts):
            if i % 2 == 1:
                # Code block
                self._add_code_block(word, part)
            else:
                self._render_text_block(word, part)

    def _render_text_block(self, word: WordDocument, block: str) -> None:
        for line in block.splitlines():
            m = _HEADING_RE.match(line)
            if m:
                level = min(len(m.group(1)), 4)
                word.add_heading(m.group(2).strip(), level=level)
                continue

            line = line.strip()
            if not line:
                continue

            if line.startswith("- ") or line.startswith("* "):
                word.add_paragraph(line[2:], style="List Bullet")
                continue

            if re.match(r"^\d+\. ", line):
                word.add_paragraph(re.sub(r"^\d+\. ", "", line), style="List Number")
                continue

            if line.startswith("> "):
                word.add_paragraph(line[2:], style="Intense Quote")
                continue

            p = word.add_paragraph()
            text = _BOLD_RE.sub(r"\1", _ITALIC_RE.sub(r"\1", line))
            p.add_run(text)

    def _add_code_block(self, word: WordDocument, code: str) -> None:
        p = word.add_paragraph()
        run = p.add_run(code.strip())
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x2D, 0x2D, 0x2D)

    def _to_bytes(self, word: WordDocument) -> bytes:
        buf = BytesIO()
        word.save(buf)
        return buf.getvalue()
